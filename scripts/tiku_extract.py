"""Read-only extraction from downloaded originals. Intermediate text stays local."""
from pathlib import Path
import hashlib,json,re
import fitz
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts'/'tiku'
OUT.mkdir(parents=True,exist_ok=True)
rows=[]
for file in sorted((ROOT/'tiku').rglob('*')):
    if file.suffix.lower() not in ('.pdf','.docx'):continue
    raw=file.read_bytes();ident='tiku-'+hashlib.sha256(str(file.relative_to(ROOT/'tiku')).encode()).hexdigest()[:12]
    pages=[];error=None
    try:
        if file.suffix.lower()=='.pdf':
            with fitz.open(file) as pdf:pages=[page.get_text(sort=True) for page in pdf]
        else:
            doc=Document(file);pages=['\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)]
    except Exception as exc:error=type(exc).__name__
    text='\n'.join(f'\n=== PAGE {i+1} ===\n{page}' for i,page in enumerate(pages))
    (OUT/f'{ident}.txt').write_text(text,encoding='utf-8')
    row=dict(id=ident,name=file.stem,subject=file.parent.name,year=(re.search(r'20\d{2}',file.stem) or [''])[0],filename=file.name,relative=str(file.relative_to(ROOT/'tiku')),sha256=hashlib.sha256(raw).hexdigest(),pages=len(pages),chars=len(text),size=len(raw),answerMarkers=len(re.findall(r'答案|正确答案|参考答案',text)),textFile=f'{ident}.txt',error=error)
    rows.append(row)
(OUT/'inventory.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(files=len(rows),pages=sum(r['pages'] for r in rows),chars=sum(r['chars'] for r in rows),withAnswers=sum(r['answerMarkers']>0 for r in rows),errors=[r['name'] for r in rows if r['error']]),ensure_ascii=False))
